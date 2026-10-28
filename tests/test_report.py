import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

import madhippu
from madhippu import cli, pdf, xlsx

from .pdfs import make_pdf

FIXTURES = Path(__file__).parent / "fixtures" / "synthetic"
BS = (FIXTURES / "indas_consolidated_balance_sheet.txt").read_text(encoding="utf-8")


@pytest.fixture
def report_pdf(tmp_path):
    # A cover with no text layer, then the balance sheet.
    path = tmp_path / "report.pdf"
    path.write_bytes(make_pdf(["", BS]))
    return path


def test_pdf_text_layer_reads_line_by_line(report_pdf):
    pages = pdf.pages(report_pdf)

    assert pages[0] == ""
    assert "Right-of-use assets 5 2,140 1,985" in [line.strip() for line in pages[1].splitlines()]


def test_parse_pdf_to_versioned_json(report_pdf):
    data = madhippu.to_dict(madhippu.parse(report_pdf))

    assert data["schema"] == "madhippu/1"
    assert "not investment advice" in data["disclaimer"].lower()
    assert data["source"]["file"] == "report.pdf" and len(data["source"]["sha256"]) == 64
    assert data["source"]["pages_without_text"] == [1]
    assert data["warnings"] == ["no text layer on page 1 (scans or images); a statement on them can't be read"]
    [st] = data["statements"]
    assert (st["kind"], st["basis"], st["page"], st["unit"]) == ("balance_sheet", "consolidated", 2, "crore")
    assert st["periods"][0] == {"start": "2024-03-31", "end": "2024-03-31"}
    assert st["warnings"] == [] and st["failed_checks"] == []
    assets = {(f["item"], f["period_end"]): f for f in st["facts"]}[("Assets", "2024-03-31")]
    assert assets == {"item": "Assets", "period_start": "2024-03-31", "period_end": "2024-03-31",
                      "value": 593_580_000_000, "unit": "INR", "decimals": -7}
    json.dumps(data)  # serialisable as is


def test_statement_without_a_unit_keeps_figures_as_printed():
    report = madhippu.parse_pages([BS])
    report.statements[0].unit = 0  # as parsed when the page states none
    report.statements[0].warnings.append("page 1: no unit (crore/lakh) stated")
    [st] = madhippu.to_dict(report)["statements"]

    assert st["unit"] == "unknown"
    assets = next(f for f in st["facts"] if f["item"] == "Assets" and f["period_end"] == "2024-03-31")
    assert (assets["value"], assets["unit"]) == (59_358, "unknown")
    assert any("no unit" in w for w in st["warnings"])


def test_no_statements_is_reported():
    report = madhippu.parse_pages(["Chairman's letter. " * 10])

    assert report.statements == [] and report.warnings == ["no financial statements found"]


def test_excel_workbook_carries_the_disclaimer_on_every_sheet(report_pdf, tmp_path):
    out = tmp_path / "out.xlsx"
    xlsx.write(madhippu.parse(report_pdf), out)

    wb = load_workbook(out)
    assert wb.sheetnames == ["About", "Consolidated BS p2"]
    about, sheet = wb["About"], wb["Consolidated BS p2"]
    assert "Not investment advice" in about["B2"].value
    assert "Not investment advice" in sheet["A1"].value
    assert sheet["A3"].value.startswith("Amounts in ₹ crore")
    rows = {r[0]: r[1:] for r in sheet.iter_rows(min_row=6, values_only=True)}
    assert rows["Assets"] == ("₹ crore", 59_358, 56_839)


def test_text_from_the_pdf_never_becomes_an_excel_formula(tmp_path):
    report = madhippu.parse_pages([])
    report.warnings = ['=HYPERLINK("http://example.invalid","x")']
    out = tmp_path / "out.xlsx"
    xlsx.write(report, out)

    cell = load_workbook(out)["About"]["B10"]
    assert cell.data_type == "s" and cell.value.startswith("=HYPERLINK")


def test_cli_writes_json_and_excel(report_pdf, tmp_path, capsys):
    status = cli.main([str(report_pdf), "--json", str(tmp_path / "o.json"), "--xlsx", str(tmp_path / "o.xlsx")])

    assert status == 0
    assert json.loads((tmp_path / "o.json").read_text())["statements"][0]["page"] == 2
    assert (tmp_path / "o.xlsx").exists()
    err = capsys.readouterr().err
    assert "page 2: consolidated balance sheet (2024-03-31, 2023-03-31)" in err
    assert "Not investment advice" in err


def test_cli_rejects_a_file_that_is_not_a_pdf(tmp_path, capsys):
    bad = tmp_path / "bad.pdf"
    bad.write_text("not a pdf")

    assert cli.main([str(bad)]) == 2
    assert "can't read" in capsys.readouterr().err
