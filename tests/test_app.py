import http.client
import json
import re
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from madhippu import app

from .pdfs import make_pdf

FIXTURES = Path(__file__).parent / "fixtures" / "synthetic"
BS = (FIXTURES / "indas_consolidated_balance_sheet.txt").read_text(encoding="utf-8")


@pytest.fixture
def server():
    server = ThreadingHTTPServer(("127.0.0.1", 0), BaseHTTPRequestHandler)
    port = server.server_address[1]
    server.RequestHandlerClass = app.make_handler(port, "secret", app.Reports(), server.shutdown)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    yield port
    server.shutdown()
    server.server_close()


def request(port, method, path, body=None, headers=None):
    conn = http.client.HTTPConnection("127.0.0.1", port, timeout=10)
    conn.request(method, path, body=body, headers=headers or {})
    res = conn.getresponse()
    return res.status, dict(res.getheaders()), res.read()


def parse(port, data, token="secret", filename="Example Report.pdf"):
    return request(port, "POST", "/parse", data,
                   {"X-Madhippu-Token": token, "X-Filename": filename})


def test_labels_read_like_report_lines():
    assert app.label("LeaseLiabilitiesNoncurrent") == "Lease liabilities non-current"
    assert app.label("ProfitLossForPeriod") == "Profit (loss) for period"
    assert app.label("CashFlowsFromUsedInOperatingActivities") == "Cash flows from (used in) operating activities"
    assert app.label("IncomeTaxesPaidRefundClassifiedAsOperatingActivities") == "Income taxes paid (net of refunds)"


def test_page_carries_this_runs_token(server):
    status, headers, body = request(server, "GET", "/")

    assert status == 200
    assert 'const TOKEN = "secret"' in body.decode()
    assert "frame-ancestors 'none'" in headers["Content-Security-Policy"]


def test_parse_shows_statements_as_printed_and_offers_downloads(server):
    status, _, body = parse(server, make_pdf(["", BS]))
    data = json.loads(body)

    assert status == 200
    assert data["file"] == "Example Report.pdf"
    [st] = data["statements"]
    assert (st["kind"], st["basis"], st["page"], st["unit"]) == ("balance_sheet", "consolidated", 2, "₹ crore")
    assert st["columns"] == ["31 Mar 2024", "31 Mar 2023"]
    assert {"item": "Assets", "label": "Assets", "values": [59_358, 56_839],
            "per_share": False, "count": False} in st["rows"]
    assert any("no text layer on page 1" in w for w in data["warnings"])

    status, headers, body = request(server, "GET", f"/report/{data['id']}.xlsx")
    assert status == 200
    assert 'filename="Example-Report-madhippu.xlsx"' in headers["Content-Disposition"]
    assert load_workbook(BytesIO(body)).sheetnames == ["About", "Consolidated BS p2"]

    status, _, body = request(server, "GET", f"/report/{data['id']}.json")
    assert json.loads(body)["schema"] == "madhippu/1"

    status, headers, body = request(server, "GET", f"/report/{data['id']}.pdf")
    assert (status, headers["Content-Type"]) == (200, "application/pdf")
    assert body.startswith(b"%PDF")


def test_refuses_posts_without_the_token(server):
    status, _, _ = parse(server, make_pdf([BS]), token="guess")

    assert status == 403


def test_refuses_other_host_names(server):
    # A site that rebinds its own name to 127.0.0.1 still sends its name.
    status, _, _ = request(server, "GET", "/", headers={"Host": f"evil.example:{server}"})

    assert status == 403


@pytest.mark.parametrize("data, message", [
    (b"not a pdf", "doesn't look like a PDF"),
    (b"%PDF-1.4 truncated", "Couldn't read that PDF"),
])
def test_explains_unreadable_files(server, data, message):
    status, _, body = parse(server, data)

    assert status == 400
    assert message in json.loads(body)["error"]


def test_unknown_report_is_not_found(server):
    status, _, _ = request(server, "GET", "/report/0123456789abcdef.xlsx")

    assert status == 404


def test_keeps_only_recent_reports():
    reports = app.Reports()
    keys = [reports.add(app.Report(f"{n}.pdf", f"{n:x}" + "0" * 63, 1, []), b"") for n in range(app.KEEP_REPORTS + 1)]

    assert reports.get(keys[0]) is None
    assert reports.get(keys[-1]) is not None


def test_downloaded_file_names_are_safe():
    assert app._safe_name('Annual "Report"; 2026.pdf') == "Annual-Report-2026"
    assert re.fullmatch(r"[\w.-]+", app._safe_name("../../etc/passwd"))
